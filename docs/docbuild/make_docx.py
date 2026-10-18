"""由同一份章节 Markdown 生成可编辑的 .docx（Word 打开即自动更新目录与页码）。

与 make_pdf.py 共用 parse()，保证 PDF 与 DOCX 的内容完全同源。
封面上的姓名 / 组号留空，方便小组确认后直接填写再导出 PDF。
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from make_pdf import MANIFEST, parse  # noqa: E402

BUILD = pathlib.Path(__file__).parent
IMG = BUILD.parent / "img"
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else BUILD / "out.docx"

SONG, HEI, KAI, MONO = "宋体", "黑体", "楷体", "Consolas"
CODE_BG = "F5F5F2"


# ---------------------------------------------------------------- 低层工具
def set_run_font(run, latin=SONG, ea=None, size=None, bold=None,
                 color=None, mono=False):
    if mono:
        latin, ea = MONO, SONG
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea or latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(paragraph, hexcolor):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def field(paragraph, instr):
    """插入一个 Word 域（TOC / PAGE 等）。"""
    r1 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    fc.set(qn("w:dirty"), "true")
    r1._element.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._element.append(it)
    r3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    r3._element.append(fc2)
    r4 = paragraph.add_run("…")
    r5 = paragraph.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    r5._element.append(fc3)
    return r4


def auto_update_fields(doc):
    """让 Word 打开文档时自动刷新目录与页码。"""
    st = doc.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    st.append(el)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(para, text, size=10.5):
    """把 **加粗** 与 `代码` 转成真实 Run。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2])
            set_run_font(r, HEI, HEI, size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1])
            set_run_font(r, MONO, SONG, size, color="A63A5B")
        else:
            r = para.add_run(part)
            set_run_font(r, SONG, SONG, size)


# ---------------------------------------------------------------- 样式初始化
def init_styles(doc):
    n = doc.styles["Normal"]
    n.font.name = SONG
    n.font.size = Pt(10.5)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    n.paragraph_format.line_spacing = 1.5
    n.paragraph_format.space_after = Pt(0)

    for name, size in (("Heading 1", 16), ("Heading 2", 14),
                       ("Heading 3", 12), ("Heading 4", 11)):
        st = doc.styles[name]
        st.font.name = HEI
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("1A2230")
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEI)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.35


# ---------------------------------------------------------------- 文档构建
def add_cover(doc):
    def blank(n=1):
        for _ in range(n):
            doc.add_paragraph()

    def centered(text, size, bold=True, font=HEI, space=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        r = p.add_run(text)
        set_run_font(r, font, font, size, bold)
        return p

    blank(4)
    centered("设计模式大作业", 30, space=26)
    centered("游戏战斗与技能系统", 20, space=8)
    centered("—— 基于 C++17 的八种设计模式综合实现", 13, bold=False, font=SONG,
             space=40)
    for line in ["院　　系：软件学院",
                 "专　　业：＿＿＿＿＿＿＿＿＿＿",
                 "组　　号：＿＿＿＿＿＿＿＿＿＿"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(2.6)
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(line), SONG, SONG, 12)
    doc.add_paragraph()
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.8)
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run("姓　　名：＿＿＿＿＿＿　班级：＿＿＿＿　"
                               "学号：＿＿＿＿＿＿"), SONG, SONG, 11.5)
    blank(4)
    centered("大连理工大学", 18, space=2)
    centered("Dalian University of Technology", 11, bold=False, font=SONG)
    doc.add_page_break()


def add_toc(doc):
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(h.add_run("目　　录"), HEI, HEI, 16, True)
    h.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph()
    field(p, r' TOC \o "1-3" \h \z \u ')
    doc.add_page_break()


def add_code(doc, src):
    for line in src.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.5)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        shade(p, CODE_BG)
        r = p.add_run(line if line else " ")
        set_run_font(r, MONO, SONG, 8)


def add_table(doc, rows):
    grid = []
    for i, row in enumerate(rows):
        if i and all(re.fullmatch(r":?-{2,}:?", c)
                     for c in row.strip().strip("|").split("|")):
            continue
        grid.append([c.strip() for c in row.strip().strip("|").split("|")])
    if not grid:
        return
    ncol = max(len(r) for r in grid)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(grid):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cp = cells[ci].paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.15
            add_inline(cp, txt, 8.6)
            if ri == 0:
                for r in cp.runs:
                    r.font.bold = True
                    set_run_font(r, HEI, HEI, 8.8, bold=True)
    doc.add_paragraph()


def add_figure(doc, name, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(IMG / name), width=Cm(15.2))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    set_run_font(c.add_run(caption), SONG, SONG, 10)


def add_blocks(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == "h":
            lvl, text = b[1], b[2]
            level = 1 if lvl <= 2 else (2 if lvl == 3 else 3)
            h = doc.add_heading(level=level)
            add_inline(h, text, {1: 16, 2: 14, 3: 12}[level])
            for r in h.runs:
                r.font.bold = True
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, b[1])
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.5)
            add_inline(p, b[1])
            for r in p.runs:
                set_run_font(r, KAI, KAI, 10.5)
        elif kind == "list":
            _, ordered, items, start = b
            for k, it in enumerate(items):
                p = doc.add_paragraph(style="List Number" if ordered
                                      else "List Bullet")
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.4
                add_inline(p, it)
        elif kind == "code":
            add_code(doc, b[1])
            doc.add_paragraph()
        elif kind == "table":
            add_table(doc, b[1])


def add_footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("— "), SONG, SONG, 9, color="555555")
    r = field(p, " PAGE ")
    set_run_font(r, SONG, SONG, 9, color="555555")
    set_run_font(p.add_run(" —"), SONG, SONG, 9, color="555555")


def build():
    doc = Document()
    init_styles(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)

    add_cover(doc)
    add_toc(doc)

    for item in MANIFEST:
        if item[0] == "md":
            add_blocks(doc, parse((BUILD / item[1]).read_text(encoding="utf-8")))
        else:
            add_figure(doc, item[1], item[2])

    add_footer(doc)
    auto_update_fields(doc)
    doc.save(str(OUT))
    print("OK ->", OUT)


if __name__ == "__main__":
    build()
